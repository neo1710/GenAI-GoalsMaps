import datetime
import os
from pathlib import Path
from typing import Any, Dict, List, Literal, Optional, Tuple

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

from app.config import settings
from app.schemas import (
    CreateWordRequest,
    CreateWordResponse,
    WordInspectResponse,
    WordKpiCard,
    WordSectionSpec,
    WordTableSpec,
)
from app.tools.file_manager import FileManager


# Color themes for Word styling
THEMES = {
    "corporate_blue": {
        "primary": RGBColor(0x1E, 0x3A, 0x8A),     # Navy #1E3A8A
        "primary_hex": "1E3A8A",
        "secondary": RGBColor(0x25, 0x63, 0xEB),   # Blue #2563EB
        "secondary_hex": "2563EB",
        "callout_bg": "EFF6FF",                    # Light blue #EFF6FF
        "zebra_bg": "F8FAFC",
        "text_dark": RGBColor(0x1E, 0x29, 0x3B),
        "text_muted": RGBColor(0x64, 0x74, 0x8B),
    },
    "emerald": {
        "primary": RGBColor(0x06, 0x5F, 0x46),     # Deep Emerald #065F46
        "primary_hex": "065F46",
        "secondary": RGBColor(0x05, 0x96, 0x69),   # Emerald #059669
        "secondary_hex": "059669",
        "callout_bg": "F0FDF4",
        "zebra_bg": "F8FAFC",
        "text_dark": RGBColor(0x1E, 0x29, 0x3B),
        "text_muted": RGBColor(0x64, 0x74, 0x8B),
    },
    "slate": {
        "primary": RGBColor(0x0F, 0x17, 0x2A),     # Slate 900
        "primary_hex": "0F172A",
        "secondary": RGBColor(0x47, 0x55, 0x69),   # Slate 600
        "secondary_hex": "475569",
        "callout_bg": "F1F5F9",
        "zebra_bg": "F8FAFC",
        "text_dark": RGBColor(0x1E, 0x29, 0x3B),
        "text_muted": RGBColor(0x64, 0x74, 0x8B),
    },
    "violet": {
        "primary": RGBColor(0x4C, 0x1D, 0x95),     # Violet 900
        "primary_hex": "4C1D95",
        "secondary": RGBColor(0x7C, 0x3A, 0xED),   # Violet 600
        "secondary_hex": "7C3AED",
        "callout_bg": "FAF5FF",
        "zebra_bg": "F8FAFC",
        "text_dark": RGBColor(0x1E, 0x29, 0x3B),
        "text_muted": RGBColor(0x64, 0x74, 0x8B),
    },
    "amber": {
        "primary": RGBColor(0x78, 0x35, 0x0F),     # Amber 900
        "primary_hex": "78350F",
        "secondary": RGBColor(0xD9, 0x77, 0x06),   # Amber 600
        "secondary_hex": "D97706",
        "callout_bg": "FFFBEB",
        "zebra_bg": "F8FAFC",
        "text_dark": RGBColor(0x1E, 0x29, 0x3B),
        "text_muted": RGBColor(0x64, 0x74, 0x8B),
    },
}


class WordTool:
    """Comprehensive tool for reading, inspecting, extracting, and creating styled Word (.docx) documents."""

    @classmethod
    def inspect_file(cls, filename: str, folder: Optional[Literal["input", "output"]] = None) -> WordInspectResponse:
        """Inspects a Word (.docx) file and returns headings, paragraphs, tables, and full markdown text."""
        path = FileManager.resolve_path(filename, folder=folder)
        if not path.exists():
            raise FileNotFoundError(f"Word document '{filename}' not found.")

        doc = docx.Document(path)
        core_props = doc.core_properties

        headings: List[Dict[str, Any]] = []
        paragraphs_text: List[str] = []
        markdown_lines: List[str] = []
        total_words = 0

        doc_title = core_props.title or ""

        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue

            words = len(text.split())
            total_words += words
            paragraphs_text.append(text)

            style_name = p.style.name.lower() if p.style else ""
            if "heading 1" in style_name:
                headings.append({"level": 1, "text": text})
                markdown_lines.append(f"\n# {text}\n")
                if not doc_title:
                    doc_title = text
            elif "heading 2" in style_name:
                headings.append({"level": 2, "text": text})
                markdown_lines.append(f"\n## {text}\n")
            elif "heading 3" in style_name:
                headings.append({"level": 3, "text": text})
                markdown_lines.append(f"\n### {text}\n")
            elif "title" in style_name or (doc_title and text == doc_title):
                doc_title = text
                markdown_lines.append(f"\n# {text}\n")
            elif "list" in style_name or p.text.startswith("- ") or p.text.startswith("* "):
                markdown_lines.append(f"- {text.lstrip('-* ')}")
            else:
                markdown_lines.append(text)

        # Extract Tables
        tables_data: List[Dict[str, Any]] = []
        for t_idx, table in enumerate(doc.tables):
            if not table.rows:
                continue

            headers = [cell.text.strip() for cell in table.rows[0].cells]
            rows_data: List[List[str]] = []
            records: List[Dict[str, str]] = []

            for row in table.rows[1:]:
                row_vals = [cell.text.strip() for cell in row.cells]
                rows_data.append(row_vals)
                rec = {headers[i] if i < len(headers) else f"Col{i}": val for i, val in enumerate(row_vals)}
                records.append(rec)

            tables_data.append({
                "table_index": t_idx + 1,
                "headers": headers,
                "row_count": len(rows_data),
                "column_count": len(headers),
                "rows": rows_data,
                "records": records,
            })

            if headers:
                markdown_lines.append("\n| " + " | ".join(headers) + " |")
                markdown_lines.append("| " + " | ".join(["---"] * len(headers)) + " |")
                for r in rows_data:
                    markdown_lines.append("| " + " | ".join(r) + " |")
                markdown_lines.append("")

        full_markdown = "\n\n".join(markdown_lines).strip()

        return WordInspectResponse(
            success=True,
            filename=path.name,
            relative_path=f"{folder or 'workspace'}/{path.name}",
            title=doc_title or path.stem,
            author=core_props.author or None,
            paragraph_count=len(paragraphs_text),
            word_count=total_words,
            table_count=len(doc.tables),
            headings=headings,
            paragraphs_preview=paragraphs_text[:10],
            tables=tables_data,
            markdown_content=full_markdown,
            file_size_bytes=path.stat().st_size,
        )

    @classmethod
    def read_text(cls, filename: str, folder: Optional[Literal["input", "output"]] = None) -> str:
        """Extracts the entire textual and tabular content of a Word file formatted as clean Markdown."""
        inspect_res = cls.inspect_file(filename, folder=folder)
        return inspect_res.markdown_content

    @classmethod
    def extract_tables(cls, filename: str, folder: Optional[Literal["input", "output"]] = None) -> List[Dict[str, Any]]:
        """Extracts all tables from a Word file as list of structured dictionaries."""
        inspect_res = cls.inspect_file(filename, folder=folder)
        return inspect_res.tables

    @classmethod
    def create_word(cls, req: CreateWordRequest) -> CreateWordResponse:
        """
        Creates a polished, styled Word document (.docx) according to structured instructions.
        """
        target_filename = req.filename
        if not target_filename.lower().endswith(".docx"):
            target_filename += ".docx"

        out_path = FileManager.get_output_path(target_filename)
        out_path.parent.mkdir(parents=True, exist_ok=True)

        doc = docx.Document()

        for section in doc.sections:
            section.top_margin = Inches(1.0)
            section.bottom_margin = Inches(1.0)
            section.left_margin = Inches(1.0)
            section.right_margin = Inches(1.0)

        theme_data = THEMES.get(req.theme or "corporate_blue", THEMES["corporate_blue"])

        core_props = doc.core_properties
        core_props.title = req.document_title or "Report"
        if req.author:
            core_props.author = req.author
        core_props.created = datetime.datetime.now(datetime.timezone.utc)

        # Title & Header Block
        if req.document_title:
            try:
                title_p = doc.add_paragraph(style="Title")
            except Exception:
                title_p = doc.add_paragraph()
            title_p.paragraph_format.space_before = Pt(0)
            title_p.paragraph_format.space_after = Pt(4)
            run = title_p.add_run(req.document_title)
            run.font.name = "Calibri"
            run.font.size = Pt(24)
            run.font.bold = True
            run.font.color.rgb = theme_data["primary"]

        if req.subtitle:
            try:
                sub_p = doc.add_paragraph(style="Subtitle")
            except Exception:
                sub_p = doc.add_paragraph()
            sub_p.paragraph_format.space_before = Pt(0)
            sub_p.paragraph_format.space_after = Pt(12)
            sub_run = sub_p.add_run(req.subtitle)
            sub_run.font.name = "Calibri"
            sub_run.font.size = Pt(13)
            sub_run.font.italic = True
            sub_run.font.color.rgb = theme_data["text_muted"]

        meta_items = []
        if req.author:
            meta_items.append(f"Author: {req.author}")
        if req.date_str:
            meta_items.append(f"Date: {req.date_str}")
        else:
            meta_items.append(f"Generated: {datetime.date.today().strftime('%B %d, %Y')}")

        if meta_items:
            meta_p = doc.add_paragraph()
            meta_p.paragraph_format.space_before = Pt(0)
            meta_p.paragraph_format.space_after = Pt(16)
            meta_run = meta_p.add_run(" • ".join(meta_items))
            meta_run.font.name = "Calibri"
            meta_run.font.size = Pt(9.5)
            meta_run.font.color.rgb = theme_data["text_muted"]

            divider_p = doc.add_paragraph()
            divider_p.paragraph_format.space_before = Pt(0)
            divider_p.paragraph_format.space_after = Pt(18)
            pBdr = parse_xml(r'<w:pBdr {}><w:bottom w:val="single" w:sz="8" w:space="1" w:color="D1D5DB"/></w:pBdr>'.format(nsdecls('w')))
            divider_p._p.get_or_add_pPr().append(pBdr)

        sections_data = req.sections or []
        if not sections_data and req.template:
            sections_data = cls._generate_template_sections(req.template, req.document_title)

        if not sections_data:
            sections_data = [
                WordSectionSpec(
                    heading="Executive Overview",
                    level=1,
                    paragraphs=["This document was automatically generated by the Sandbox Service agent."],
                    callout="Key Takeaway: Automated report generation is operational.",
                )
            ]

        total_words = len(req.document_title.split()) if req.document_title else 0
        total_tables = 0
        markdown_summary_parts: List[str] = [f"# {req.document_title or 'Document'}"]

        for sec in sections_data:
            if sec.heading:
                head_p = doc.add_heading(sec.heading, level=min(sec.level, 3))
                head_p.paragraph_format.space_before = Pt(16)
                head_p.paragraph_format.space_after = Pt(6)
                head_p.paragraph_format.keep_with_next = True
                
                for r in head_p.runs:
                    r.font.name = "Calibri"
                    if sec.level == 1:
                        r.font.color.rgb = theme_data["primary"]
                    elif sec.level == 2:
                        r.font.color.rgb = theme_data["secondary"]
                    else:
                        r.font.color.rgb = theme_data["text_dark"]

                if sec.level == 1:
                    markdown_summary_parts.append(f"\n## {sec.heading}")
                elif sec.level == 2:
                    markdown_summary_parts.append(f"\n### {sec.heading}")
                else:
                    markdown_summary_parts.append(f"\n#### {sec.heading}")

            if sec.kpis:
                total_tables += 1
                cls._render_kpi_cards(doc, sec.kpis, theme_data)

            if sec.callout:
                cls._render_callout_box(doc, sec.callout, theme_data)
                markdown_summary_parts.append(f"> **Highlight:** {sec.callout}")

            if sec.paragraphs:
                for para_text in sec.paragraphs:
                    p = doc.add_paragraph()
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(8)
                    p.paragraph_format.line_spacing = 1.15
                    run = p.add_run(para_text)
                    run.font.name = "Calibri"
                    run.font.size = Pt(11)
                    run.font.color.rgb = theme_data["text_dark"]
                    total_words += len(para_text.split())
                    markdown_summary_parts.append(para_text)

            if sec.bullet_points:
                for bp in sec.bullet_points:
                    bp_para = doc.add_paragraph(style="List Bullet")
                    bp_para.paragraph_format.space_before = Pt(0)
                    bp_para.paragraph_format.space_after = Pt(3)
                    run = bp_para.add_run(bp)
                    run.font.name = "Calibri"
                    run.font.size = Pt(11)
                    run.font.color.rgb = theme_data["text_dark"]
                    total_words += len(bp.split())
                    markdown_summary_parts.append(f"- {bp}")

            if sec.numbered_list:
                for idx, item in enumerate(sec.numbered_list, start=1):
                    num_para = doc.add_paragraph(style="List Number")
                    num_para.paragraph_format.space_before = Pt(0)
                    num_para.paragraph_format.space_after = Pt(3)
                    run = num_para.add_run(item)
                    run.font.name = "Calibri"
                    run.font.size = Pt(11)
                    run.font.color.rgb = theme_data["text_dark"]
                    total_words += len(item.split())
                    markdown_summary_parts.append(f"{idx}. {item}")

            if sec.table and sec.table.headers:
                total_tables += 1
                cls._render_table(doc, sec.table, theme_data)
                markdown_summary_parts.append(f"[Table: {len(sec.table.rows)} rows]")

        doc.save(out_path)

        return CreateWordResponse(
            success=True,
            filename=out_path.name,
            relative_path=f"output/{out_path.name}",
            absolute_path=str(out_path),
            download_url=f"/files/download/output/{out_path.name}",
            word_count=total_words,
            section_count=len(sections_data),
            table_count=total_tables,
            file_size_bytes=out_path.stat().st_size,
            markdown_summary="\n\n".join(markdown_summary_parts),
            message=f"Created professional Word document '{out_path.name}' with {len(sections_data)} sections, {total_tables} tables, and ~{total_words} words.",
        )

    @classmethod
    def _render_kpi_cards(cls, doc: docx.Document, kpis: List[WordKpiCard], theme: Dict[str, Any]):
        """Renders side-by-side KPI metric cards inside a styled table."""
        count = len(kpis)
        if count == 0:
            return

        table = doc.add_table(rows=1, cols=count)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True

        for idx, kpi in enumerate(kpis):
            cell = table.cell(0, idx)
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{theme["callout_bg"]}"/>')
            cell._tc.get_or_add_tcPr().append(shd)

            tcBorders = parse_xml(f'''
                <w:tcBorders {nsdecls("w")}>
                    <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                    <w:left w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                    <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                    <w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                </w:tcBorders>
            ''')
            cell._tc.get_or_add_tcPr().append(tcBorders)

            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)

            m_run = p.add_run(kpi.metric.upper() + "\n")
            m_run.font.name = "Calibri"
            m_run.font.size = Pt(8.5)
            m_run.font.bold = True
            m_run.font.color.rgb = theme["text_muted"]

            v_run = p.add_run(kpi.value + "\n")
            v_run.font.name = "Calibri"
            v_run.font.size = Pt(18)
            v_run.font.bold = True
            v_run.font.color.rgb = theme["primary"]

            if kpi.subtitle:
                s_run = p.add_run(kpi.subtitle)
                s_run.font.name = "Calibri"
                s_run.font.size = Pt(8.5)
                s_run.font.italic = True
                s_run.font.color.rgb = theme["text_muted"]

        sp = doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(0)
        sp.paragraph_format.space_after = Pt(6)

    @classmethod
    def _render_callout_box(cls, doc: docx.Document, text: str, theme: Dict[str, Any]):
        """Renders an executive highlight callout box with thick left accent bar."""
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        cell = table.cell(0, 0)
        cell.width = Inches(6.5)

        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{theme["callout_bg"]}"/>')
        cell._tc.get_or_add_tcPr().append(shd)

        borders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="none"/>
                <w:left w:val="single" w:sz="24" w:space="0" w:color="{theme["primary_hex"]}"/>
                <w:bottom w:val="none"/>
                <w:right w:val="none"/>
            </w:tcBorders>
        ''')
        cell._tc.get_or_add_tcPr().append(borders)

        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.left_indent = Pt(12)
        p.paragraph_format.right_indent = Pt(12)

        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(10.5)
        run.font.italic = True
        run.font.color.rgb = theme["text_dark"]

        sp = doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(0)
        sp.paragraph_format.space_after = Pt(6)

    @classmethod
    def _render_table(cls, doc: docx.Document, table_spec: WordTableSpec, theme: Dict[str, Any]):
        """Renders a styled table with colored header row, alternating row shading, and thin borders."""
        col_count = len(table_spec.headers)
        row_count = len(table_spec.rows) + 1

        table = doc.add_table(rows=row_count, cols=col_count)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        header_row = table.rows[0]
        for c_idx, head_text in enumerate(table_spec.headers):
            cell = header_row.cells[c_idx]
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{theme["primary_hex"]}"/>')
            cell._tc.get_or_add_tcPr().append(shd)

            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(5)

            run = p.add_run(head_text)
            run.font.name = "Calibri"
            run.font.size = Pt(10)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        zebra_fill_hex = theme["zebra_bg"]
        for r_idx, row_values in enumerate(table_spec.rows, start=1):
            row = table.rows[r_idx]
            is_zebra = (r_idx % 2 == 0)

            for c_idx, val in enumerate(row_values):
                if c_idx >= col_count:
                    break
                cell = row.cells[c_idx]

                if is_zebra:
                    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{zebra_fill_hex}"/>')
                    cell._tc.get_or_add_tcPr().append(shd)

                borders = parse_xml(r'''
                    <w:tcBorders {}>
                        <w:top w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
                        <w:bottom w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
                        <w:left w:val="none"/>
                        <w:right w:val="none"/>
                    </w:tcBorders>
                '''.format(nsdecls('w')))
                cell._tc.get_or_add_tcPr().append(borders)

                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)

                val_str = str(val if val is not None else "")
                run = p.add_run(val_str)
                run.font.name = "Calibri"
                run.font.size = Pt(9.5)
                run.font.color.rgb = theme["text_dark"]

                if table_spec.column_alignments and c_idx < len(table_spec.column_alignments):
                    align_str = table_spec.column_alignments[c_idx].lower()
                    if align_str == "right":
                        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    elif align_str == "center":
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    else:
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                else:
                    if val_str.replace(".", "", 1).replace("%", "").replace("$", "").replace(",", "").isdigit():
                        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    else:
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

        sp = doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(0)
        sp.paragraph_format.space_after = Pt(10)

    @classmethod
    def _generate_template_sections(cls, template: str, doc_title: Optional[str] = None) -> List[WordSectionSpec]:
        """Generates rich pre-built template sections for common enterprise reports."""
        if template == "executive_summary":
            return [
                WordSectionSpec(
                    heading="1. Strategic Executive Summary",
                    level=1,
                    paragraphs=[
                        "This quarterly evaluation assesses core operational objectives, performance benchmarks, and delivery trajectories across key mission workstreams. All critical milestones have progressed under active monitoring with zero catastrophic blockers.",
                        "Resource allocations and strategic initiatives remain within target variance, enabling our team to accelerate high-impact capabilities."
                    ],
                    kpis=[
                        WordKpiCard(metric="Milestones Met", value="88%", subtitle="+6% vs prior quarter"),
                        WordKpiCard(metric="Sprint Velocity", value="94 pts", subtitle="Consistent cadence"),
                        WordKpiCard(metric="Open Risks", value="2 Low", subtitle="Zero critical escalations"),
                        WordKpiCard(metric="Budget Efficiency", value="97.4%", subtitle="On track"),
                    ],
                    callout="Executive Takeaway: Operational momentum is strong. Transitioning to Phase 2 automation will compress deliverable timelines by an estimated 25%."
                ),
                WordSectionSpec(
                    heading="2. Workstream Status & Key Milestones",
                    level=1,
                    paragraphs=[
                        "Below is the consolidated breakdown of active workstreams and their delivery health:"
                    ],
                    table=WordTableSpec(
                        headers=["Initiative", "Workstream Lead", "Target Date", "Status", "Progress"],
                        rows=[
                            ["Core Infrastructure Hardening", "DevOps Team", "Nov 15, 2026", "Completed", "100%"],
                            ["Automated Data Pipelines", "Data Eng Team", "Dec 01, 2026", "In Progress", "78%"],
                            ["AI Agent Integration Layer", "GenAI Architecture", "Dec 15, 2026", "In Progress", "65%"],
                            ["Enterprise Audit & Compliance", "SecOps", "Jan 10, 2027", "Planning", "25%"],
                        ],
                        column_alignments=["left", "left", "center", "center", "right"],
                    )
                ),
                WordSectionSpec(
                    heading="3. Key Action Items & Next Steps",
                    level=1,
                    numbered_list=[
                        "Finalize end-to-end integration tests for multi-stream pipelines.",
                        "Deploy automated security validation gates into the CI/CD pipeline.",
                        "Publish the complete developer onboarding and architectural reference guide.",
                        "Schedule executive stakeholder demo for milestone validation.",
                    ]
                )
            ]

        elif template == "goal_progress_report":
            return [
                WordSectionSpec(
                    heading="1. Goal Health & Milestone Scorecard",
                    level=1,
                    paragraphs=[
                        "Comprehensive quarterly scorecard tracking progress against strategic objectives, milestone completions, and departmental accountability."
                    ],
                    kpis=[
                        WordKpiCard(metric="Total Goals", value="16", subtitle="Active across 4 teams"),
                        WordKpiCard(metric="On Schedule", value="13", subtitle="81% green status"),
                        WordKpiCard(metric="At Risk", value="3", subtitle="Mitigations assigned"),
                    ],
                    callout="Immediate Attention: Goal GOAL-004 requires cross-functional backend pairing to resolve downstream schema synchronization."
                ),
                WordSectionSpec(
                    heading="2. Detailed Goal Breakdown",
                    level=1,
                    table=WordTableSpec(
                        headers=["Goal ID", "Title", "Owner", "Priority", "Status", "% Complete"],
                        rows=[
                            ["GOAL-001", "Real-Time Agent Streaming", "Backend Eng", "Critical", "Active", "90%"],
                            ["GOAL-002", "Office Document Tooling", "Sandbox Lead", "High", "Active", "85%"],
                            ["GOAL-003", "Instructional Skill Registry", "AI Core", "High", "Active", "80%"],
                            ["GOAL-004", "External Cloud Storage Sync", "Cloud Team", "Medium", "Blocked", "35%"],
                        ],
                        column_alignments=["left", "left", "left", "center", "center", "right"],
                    )
                ),
                WordSectionSpec(
                    heading="3. Risk Mitigation Strategy",
                    level=1,
                    bullet_points=[
                        "Allocate senior engineering capacity to unblock synchronization interfaces.",
                        "Implement automated health pinging and retry timeouts on network requests.",
                        "Conduct weekly milestone reviews with stream owners until all goals turn green.",
                    ]
                )
            ]

        elif template == "technical_spec":
            return [
                WordSectionSpec(
                    heading="1. System Architecture & Objectives",
                    level=1,
                    paragraphs=[
                        "This technical specification outlines the design, API contracts, streaming protocols, and security guardrails for the agent execution engine."
                    ],
                    callout="Design Principle: Zero synchronous blocking for long-running tasks. Real-time Server-Sent Events (SSE) provide total execution visibility."
                ),
                WordSectionSpec(
                    heading="2. Component Interaction Matrix",
                    level=1,
                    table=WordTableSpec(
                        headers=["Component", "Responsibility", "Protocol", "Data Format"],
                        rows=[
                            ["API Gateway / Orchestrator", "Route requests & dispatch agent actions", "HTTP / REST", "JSON"],
                            ["Sandbox Runner", "Execute untrusted code in isolated process", "Subprocess Pipes", "Streams (SSE)"],
                            ["Office Document Tool", "Generate styled Word & Excel workbooks", "In-Process Engine", "OpenXML / Docx"],
                            ["Skill Registry", "Manage executable task recipes and schemas", "In-Memory / JSON", "Pydantic Models"],
                        ],
                        column_alignments=["left", "left", "center", "center"],
                    )
                ),
                WordSectionSpec(
                    heading="3. Security & Isolation Invariants",
                    level=1,
                    bullet_points=[
                        "Static AST and regular expression security audits prevent dangerous system and subprocess calls.",
                        "Dedicated workspace directory confinement blocks path traversal attacks.",
                        "Strict execution timeouts prevent runaway loops or denial of service.",
                    ]
                )
            ]

        return [
            WordSectionSpec(
                heading="Overview",
                level=1,
                paragraphs=[f"Report: {doc_title or 'Automated Analysis'}"],
            )
        ]

