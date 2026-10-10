import os
from pathlib import Path
from typing import Any, Dict, List, Literal, Optional, Tuple, Union
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import numpy as np

from app.config import settings
from app.schemas import (
    AnalyzeCsvResponse,
    CreateExcelRequest,
    CreateExcelResponse,
    ExcelInspectResponse,
    ExcelSheetInfo,
    ExcelSheetSpec,
)
from app.tools.csv_analyser import CSVAnalyser
from app.tools.file_manager import FileManager


# Color palettes for professional Excel generation
PALETTES = {
    "corporate_blue": {
        "header_fill": "1E3A8A",      # Deep Navy
        "header_font": "FFFFFF",
        "zebra_fill": "F1F5F9",       # Very light slate
        "total_fill": "E2E8F0",       # Light slate
        "accent": "2563EB",
    },
    "emerald": {
        "header_fill": "065F46",      # Deep Emerald
        "header_font": "FFFFFF",
        "zebra_fill": "F0FDF4",       # Very light emerald
        "total_fill": "DCFCE7",
        "accent": "059669",
    },
    "slate": {
        "header_fill": "1E293B",      # Slate 800
        "header_font": "FFFFFF",
        "zebra_fill": "F8FAFC",
        "total_fill": "E2E8F0",
        "accent": "475569",
    },
    "violet": {
        "header_fill": "4C1D95",      # Deep Violet
        "header_font": "FFFFFF",
        "zebra_fill": "FAF5FF",
        "total_fill": "EDE9FE",
        "accent": "7C3AED",
    },
    "amber": {
        "header_fill": "78350F",      # Amber/Bronze
        "header_font": "FFFFFF",
        "zebra_fill": "FFFBEB",
        "total_fill": "FEF3C7",
        "accent": "D97706",
    },
}

THIN_BORDER = Border(
    left=Side(style="thin", color="D1D5DB"),
    right=Side(style="thin", color="D1D5DB"),
    top=Side(style="thin", color="D1D5DB"),
    bottom=Side(style="thin", color="D1D5DB"),
)

DOUBLE_BOTTOM_BORDER = Border(
    left=Side(style="thin", color="D1D5DB"),
    right=Side(style="thin", color="D1D5DB"),
    top=Side(style="thin", color="D1D5DB"),
    bottom=Side(style="double", color="374151"),
)


class ExcelTool:
    """Comprehensive tool for inspecting, analyzing, and generating professional Excel files."""

    @classmethod
    def inspect_file(cls, filename: str, folder: Optional[Literal["input", "output"]] = None) -> ExcelInspectResponse:
        """Inspects an Excel (.xlsx, .xls) file: sheets, dimensions, headers, preview."""
        path = FileManager.resolve_path(filename, folder=folder)
        if not path.exists():
            raise FileNotFoundError(f"Excel file '{filename}' not found.")

        sheets_info: List[ExcelSheetInfo] = []
        with pd.ExcelFile(path) as excel_file:
            sheet_names = excel_file.sheet_names
            for sheet_name in sheet_names:
                # Detect header row (skips title banners if present)
                raw_df = excel_file.parse(sheet_name=sheet_name, header=None, nrows=10)
                header_idx = 0
                for r in range(min(len(raw_df), 10)):
                    if len(raw_df.iloc[r].dropna()) > 1:
                        header_idx = r
                        break

                df = excel_file.parse(sheet_name=sheet_name, header=header_idx, nrows=settings.MAX_CSV_PREVIEW_ROWS)
                full_df = excel_file.parse(sheet_name=sheet_name, header=header_idx)

                clean_records: List[Dict[str, Any]] = []
                for rec in df.head(10).to_dict(orient="records"):
                    clean_rec = {}
                    for k, v in rec.items():
                        if pd.isna(v):
                            clean_rec[str(k)] = None
                        elif isinstance(v, (pd.Timestamp, pd.Timedelta)):
                            clean_rec[str(k)] = str(v)
                        else:
                            clean_rec[str(k)] = v
                    clean_records.append(clean_rec)

                sheets_info.append(
                    ExcelSheetInfo(
                        sheet_name=sheet_name,
                        row_count=len(full_df),
                        column_count=len(full_df.columns),
                        columns=[str(c) for c in full_df.columns],
                        preview=clean_records,
                    )
                )

        return ExcelInspectResponse(
            success=True,
            filename=path.name,
            relative_path=f"{folder or 'workspace'}/{path.name}",
            sheet_count=len(sheet_names),
            sheet_names=sheet_names,
            sheets=sheets_info,
            file_size_bytes=path.stat().st_size,
        )

    @classmethod
    def analyze_sheet(
        cls,
        filename: str,
        sheet_name: Optional[str] = None,
        folder: Optional[Literal["input", "output"]] = None,
        generate_markdown: bool = True,
    ) -> AnalyzeCsvResponse:
        """Runs statistical profiling and anomaly analysis on a specific Excel sheet."""
        path = FileManager.resolve_path(filename, folder=folder)
        if not path.exists():
            raise FileNotFoundError(f"Excel file '{filename}' not found.")

        with pd.ExcelFile(path) as excel_file:
            target_sheet = sheet_name or excel_file.sheet_names[0]
            raw_df = excel_file.parse(sheet_name=target_sheet, header=None, nrows=10)
            header_idx = 0
            for r in range(min(len(raw_df), 10)):
                if len(raw_df.iloc[r].dropna()) > 1:
                    header_idx = r
                    break
            df = excel_file.parse(sheet_name=target_sheet, header=header_idx)

        # Convert dataframe to CSV string and run CSVAnalyser
        csv_str = df.to_csv(index=False)
        from app.schemas import AnalyzeCsvRequest
        req = AnalyzeCsvRequest(
            filename=None,
            csv_content=csv_str,
            generate_markdown_report=generate_markdown,
        )
        res = CSVAnalyser.analyze(req)
        res.source = f"{path.name} [Sheet: {target_sheet}]"
        return res

    @classmethod
    def convert_to_csv(
        cls,
        filename: str,
        sheet_name: Optional[str] = None,
        output_filename: Optional[str] = None,
        folder: Optional[Literal["input", "output"]] = None,
    ) -> Dict[str, Any]:
        """Converts an Excel sheet into a standard CSV file in workspace output/."""
        path = FileManager.resolve_path(filename, folder=folder)
        if not path.exists():
            raise FileNotFoundError(f"Excel file '{filename}' not found.")

        with pd.ExcelFile(path) as excel_file:
            target_sheet = sheet_name or excel_file.sheet_names[0]
            raw_df = excel_file.parse(sheet_name=target_sheet, header=None, nrows=10)
            header_idx = 0
            for r in range(min(len(raw_df), 10)):
                if len(raw_df.iloc[r].dropna()) > 1:
                    header_idx = r
                    break
            df = excel_file.parse(sheet_name=target_sheet, header=header_idx)

        out_name = output_filename or f"{Path(filename).stem}_{target_sheet.lower().replace(' ', '_')}.csv"
        out_path = FileManager.get_output_path(out_name)
        df.to_csv(out_path, index=False)

        return {
            "success": True,
            "source_excel": path.name,
            "sheet_name": target_sheet,
            "output_csv": out_name,
            "relative_path": f"output/{out_name}",
            "row_count": len(df),
            "column_count": len(df.columns),
            "download_url": f"/files/download/output/{out_name}",
        }

    @classmethod
    def create_excel(cls, req: CreateExcelRequest) -> CreateExcelResponse:
        """
        Creates a styled, multi-sheet Excel file (.xlsx) according to instructions and specifications.
        """
        target_filename = req.filename
        if not target_filename.lower().endswith(".xlsx"):
            target_filename += ".xlsx"

        out_path = FileManager.get_output_path(target_filename)
        out_path.parent.mkdir(parents=True, exist_ok=True)

        wb = openpyxl.Workbook()
        default_sheet = wb.active

        palette = PALETTES.get(req.theme or "corporate_blue", PALETTES["corporate_blue"])
        sheets_data = req.sheets or []

        # If no sheets provided but template requested, use built-in template
        if not sheets_data and req.template:
            sheets_data = cls._generate_template_sheets(req.template, req.template_row_count or 20)

        # If still empty, create default sample sheet
        if not sheets_data:
            sheets_data = [
                ExcelSheetSpec(
                    title="Summary",
                    columns=["Metric", "Value", "Status", "Notes"],
                    rows=[
                        ["Total Tasks", 45, "Active", "Tracking well"],
                        ["Completed", 32, "Done", "Sprint 1 & 2"],
                        ["Pending", 13, "In Progress", "Sprint 3"],
                    ],
                )
            ]

        sheets_created: List[str] = []
        total_rows = 0
        previews: Dict[str, List[Dict[str, Any]]] = {}

        for i, sheet_spec in enumerate(sheets_data):
            ws = wb.create_sheet(title=sheet_spec.title) if i > 0 or default_sheet is None else default_sheet
            ws.title = sheet_spec.title
            sheets_created.append(sheet_spec.title)

            # Enable grid lines explicitly
            ws.views.sheetView[0].showGridLines = True

            # Styling definitions
            header_fill_color = sheet_spec.header_bg_color or palette["header_fill"]
            header_font_color = sheet_spec.header_text_color or palette["header_font"]

            header_fill = PatternFill(start_color=header_fill_color, end_color=header_fill_color, fill_type="solid")
            header_font = Font(name="Calibri", size=11, bold=True, color=header_font_color)
            header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            zebra_fill = PatternFill(start_color=palette["zebra_fill"], end_color=palette["zebra_fill"], fill_type="solid")
            total_fill = PatternFill(start_color=palette["total_fill"], end_color=palette["total_fill"], fill_type="solid")
            bold_font = Font(name="Calibri", size=11, bold=True)
            regular_font = Font(name="Calibri", size=11)

            # Optional Title Banner row
            current_row = 1
            if req.document_title and i == 0:
                ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(sheet_spec.columns), 1))
                title_cell = ws.cell(row=1, column=1, value=req.document_title)
                title_cell.font = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
                title_cell.alignment = Alignment(horizontal="left", vertical="center")
                ws.row_dimensions[1].height = 36
                current_row = 3  # Leave blank row

            # Write Column Headers
            header_row_idx = current_row
            ws.row_dimensions[header_row_idx].height = 28
            for col_idx, col_name in enumerate(sheet_spec.columns, start=1):
                cell = ws.cell(row=header_row_idx, column=col_idx, value=col_name)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = header_alignment
                cell.border = THIN_BORDER

            # Normalize data: dict records or 2D matrix
            data_rows: List[List[Any]] = []
            if sheet_spec.rows:
                data_rows = sheet_spec.rows
            elif sheet_spec.data:
                for record in sheet_spec.data:
                    row_vals = [record.get(col, "") for col in sheet_spec.columns]
                    data_rows.append(row_vals)

            total_rows += len(data_rows)
            preview_records: List[Dict[str, Any]] = []

            # Write Data Rows
            start_data_row = header_row_idx + 1
            for r_idx, row_values in enumerate(data_rows, start=start_data_row):
                ws.row_dimensions[r_idx].height = 20
                is_zebra = (r_idx % 2 == 0) and sheet_spec.zebra_stripes

                row_dict: Dict[str, Any] = {}
                for c_idx, val in enumerate(row_values, start=1):
                    col_name = sheet_spec.columns[c_idx - 1] if c_idx - 1 < len(sheet_spec.columns) else f"Col{c_idx}"
                    cell = ws.cell(row=r_idx, column=c_idx, value=val)
                    cell.font = regular_font
                    cell.border = THIN_BORDER
                    if is_zebra:
                        cell.fill = zebra_fill

                    # Format application
                    fmt = sheet_spec.column_formats.get(col_name) if sheet_spec.column_formats else None
                    if fmt:
                        cell.number_format = cls._resolve_number_format(fmt)
                    elif isinstance(val, (int, float)):
                        cell.alignment = Alignment(horizontal="right", vertical="center")

                    if len(preview_records) < 10:
                        row_dict[col_name] = val

                if len(preview_records) < 10:
                    preview_records.append(row_dict)

            # Optional Summary / Formula Row
            last_data_row = start_data_row + len(data_rows) - 1
            if sheet_spec.summary_row and len(data_rows) > 0:
                summary_row_idx = last_data_row + 1
                ws.row_dimensions[summary_row_idx].height = 24

                for c_idx, col_name in enumerate(sheet_spec.columns, start=1):
                    cell = ws.cell(row=summary_row_idx, column=c_idx)
                    cell.font = bold_font
                    cell.fill = total_fill
                    cell.border = DOUBLE_BOTTOM_BORDER

                    col_letter = get_column_letter(c_idx)
                    formula_type = sheet_spec.summary_row.get(col_name)

                    if formula_type:
                        agg = formula_type.upper().strip()
                        cell.value = f"={agg}({col_letter}{start_data_row}:{col_letter}{last_data_row})"
                        fmt = sheet_spec.column_formats.get(col_name) if sheet_spec.column_formats else None
                        if fmt:
                            cell.number_format = cls._resolve_number_format(fmt)
                    elif c_idx == 1:
                        cell.value = "Total / Summary"
                        cell.alignment = Alignment(horizontal="left", vertical="center")

            # Auto-fit Column Widths
            for col in ws.columns:
                col_letter = get_column_letter(col[0].column)
                specified_width = (sheet_spec.column_widths or {}).get(sheet_spec.columns[col[0].column - 1]) if col[0].column - 1 < len(sheet_spec.columns) else None
                if specified_width:
                    ws.column_dimensions[col_letter].width = max(specified_width, 10)
                else:
                    max_len = 0
                    for cell in col:
                        if cell.row == 1 and req.document_title:
                            continue
                        val_str = str(cell.value or "")
                        max_len = max(max_len, len(val_str))
                    ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

            previews[sheet_spec.title] = preview_records

        wb.save(out_path)

        return CreateExcelResponse(
            success=True,
            filename=out_path.name,
            relative_path=f"output/{out_path.name}",
            absolute_path=str(out_path),
            download_url=f"/files/download/output/{out_path.name}",
            sheets_created=sheets_created,
            total_rows=total_rows,
            file_size_bytes=out_path.stat().st_size,
            previews=previews,
            message=f"Created professional Excel workbook '{out_path.name}' with {len(sheets_created)} sheet(s) and {total_rows} total rows.",
        )

    @classmethod
    def _resolve_number_format(cls, fmt: str) -> str:
        fmt_lower = fmt.lower().strip()
        if fmt_lower in ("currency", "usd", "$"):
            return "$#,##0.00"
        if fmt_lower in ("currency_integer", "$_int"):
            return "$#,##0"
        if fmt_lower in ("percent", "percentage", "%"):
            return "0.0%"
        if fmt_lower in ("percent_int", "%_int"):
            return "0%"
        if fmt_lower in ("integer", "int", "number"):
            return "#,##0"
        if fmt_lower in ("float", "decimal"):
            return "#,##0.00"
        if fmt_lower in ("date", "yyyy-mm-dd"):
            return "yyyy-mm-dd"
        return fmt

    @classmethod
    def _generate_template_sheets(cls, template: str, row_count: int = 20) -> List[ExcelSheetSpec]:
        np.random.seed(42)

        if template == "kpi_dashboard":
            statuses = ["Completed", "In Progress", "At Risk", "Planning"]
            categories = ["Engineering", "Product", "Growth", "Customer Success"]
            owners = ["Sarah Chen", "Marcus Vance", "Elena Rostova", "Devon Lee"]
            rows = []
            for i in range(1, row_count + 1):
                cat = str(np.random.choice(categories))
                status = str(np.random.choice(statuses, p=[0.4, 0.35, 0.15, 0.1]))
                pct = 100 if status == "Completed" else np.random.randint(10, 85)
                budget = np.random.randint(5000, 75000)
                spent = int(budget * (pct / 100.0) * np.random.uniform(0.8, 1.1))
                rows.append([
                    f"GOAL-{i:03d}",
                    f"{cat} Milestone Phase {i}",
                    cat,
                    owners[i % len(owners)],
                    status,
                    pct / 100.0,
                    budget,
                    spent,
                    budget - spent,
                ])

            return [
                ExcelSheetSpec(
                    title="Goals_and_KPIs",
                    columns=["Goal ID", "Milestone Title", "Department", "Owner", "Status", "Progress", "Budget USD", "Spent USD", "Remaining USD"],
                    rows=rows,
                    header_bg_color="1E3A8A",
                    column_formats={
                        "Progress": "percent",
                        "Budget USD": "currency_integer",
                        "Spent USD": "currency_integer",
                        "Remaining USD": "currency_integer",
                    },
                    summary_row={
                        "Budget USD": "SUM",
                        "Spent USD": "SUM",
                        "Remaining USD": "SUM",
                        "Progress": "AVERAGE",
                    },
                    zebra_stripes=True,
                )
            ]

        elif template == "financial_model":
            departments = ["Engineering", "Marketing", "Operations", "Sales", "G&A"]
            categories = ["Software Licenses", "Personnel", "Cloud Infrastructure", "Events & Travel", "Consulting"]
            rows = []
            for i in range(1, row_count + 1):
                dept = departments[i % len(departments)]
                cat = categories[i % len(categories)]
                allocated = np.random.randint(15000, 120000)
                actual = int(allocated * np.random.uniform(0.7, 1.15))
                variance = allocated - actual
                pct_spent = actual / allocated if allocated else 0
                rows.append([
                    f"BUDGET-{i:03d}",
                    dept,
                    cat,
                    allocated,
                    actual,
                    variance,
                    pct_spent,
                    "Under Budget" if variance >= 0 else "Over Budget",
                ])

            return [
                ExcelSheetSpec(
                    title="Budget_Overview",
                    columns=["Line Item", "Department", "Expense Category", "Allocated Budget", "Actual Spend", "Variance", "% Spent", "Status"],
                    rows=rows,
                    header_bg_color="065F46",
                    column_formats={
                        "Allocated Budget": "currency_integer",
                        "Actual Spend": "currency_integer",
                        "Variance": "currency_integer",
                        "% Spent": "percent",
                    },
                    summary_row={
                        "Allocated Budget": "SUM",
                        "Actual Spend": "SUM",
                        "Variance": "SUM",
                        "% Spent": "AVERAGE",
                    },
                    zebra_stripes=True,
                )
            ]

        elif template == "project_tracker":
            sprints = ["Sprint 42", "Sprint 43", "Sprint 44"]
            task_types = ["Feature", "Bug Fix", "Refactor", "Security", "Documentation"]
            priorities = ["P0 - Blocker", "P1 - High", "P2 - Medium", "P3 - Low"]
            rows = []
            for i in range(1, row_count + 1):
                story_points = int(np.random.choice([1, 2, 3, 5, 8, 13]))
                hours_logged = round(story_points * np.random.uniform(3.5, 6.0), 1)
                rows.append([
                    f"TASK-{i:04d}",
                    sprints[i % len(sprints)],
                    f"System Optimization & Implementation #{i}",
                    task_types[i % len(task_types)],
                    priorities[i % len(priorities)],
                    story_points,
                    hours_logged,
                    "Done" if i % 3 == 0 else "In Progress",
                ])

            return [
                ExcelSheetSpec(
                    title="Sprint_Tasks",
                    columns=["Task ID", "Sprint", "Task Summary", "Type", "Priority", "Story Points", "Hours Logged", "Status"],
                    rows=rows,
                    header_bg_color="1E293B",
                    column_formats={
                        "Story Points": "integer",
                        "Hours Logged": "float",
                    },
                    summary_row={
                        "Story Points": "SUM",
                        "Hours Logged": "SUM",
                    },
                    zebra_stripes=True,
                )
            ]

        elif template == "sales_summary":
            regions = ["North America", "EMEA", "APAC", "LATAM"]
            products = ["Enterprise Cloud", "Team Workspace", "Security Suite", "API Add-on"]
            rows = []
            for i in range(1, row_count + 1):
                units = np.random.randint(5, 120)
                price = int(np.random.choice([250, 500, 1200, 3500]))
                discount = round(float(np.random.choice([0.0, 0.05, 0.1, 0.15, 0.2])), 2)
                gross = units * price
                net = gross * (1 - discount)
                rows.append([
                    f"INV-2026-{i:04d}",
                    regions[i % len(regions)],
                    products[i % len(products)],
                    units,
                    price,
                    discount,
                    gross,
                    net,
                ])

            return [
                ExcelSheetSpec(
                    title="Sales_Performance",
                    columns=["Invoice ID", "Region", "Product Tier", "Units Sold", "Unit Price", "Discount %", "Gross Revenue", "Net Revenue"],
                    rows=rows,
                    header_bg_color="78350F",
                    column_formats={
                        "Units Sold": "integer",
                        "Unit Price": "currency_integer",
                        "Discount %": "percent",
                        "Gross Revenue": "currency_integer",
                        "Net Revenue": "currency_integer",
                    },
                    summary_row={
                        "Units Sold": "SUM",
                        "Gross Revenue": "SUM",
                        "Net Revenue": "SUM",
                    },
                    zebra_stripes=True,
                )
            ]

        return [
            ExcelSheetSpec(
                title="Data",
                columns=["ID", "Name", "Score"],
                rows=[[f"ROW-{i}", f"Item {i}", np.random.randint(10, 100)] for i in range(1, row_count + 1)],
                zebra_stripes=True,
            )
        ]

